import io
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import zipfile
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))


def parse(path):
    units=[];remaining=65536;truncated=False
    def add(unit,number,text,source,**extra):
        nonlocal remaining,truncated
        if len(units)>=200 or remaining<=0:truncated=True;return
        if len(text)>remaining:truncated=True
        units.append({'unit':unit,'number':number,'source':source,'text':text[:remaining],**extra});remaining-=len(text[:remaining])
    suffix=path.suffix
    if suffix=='.pdf':
        import pdfplumber
        with pdfplumber.open(path) as document:
            for number,page in enumerate(document.pages[:200],1):
                text=page.extract_text() or '';ocr=False;needs_ocr=False
                if not text.strip():
                    command=shutil.which(os.environ.get('AGENT_TESSERACT_COMMAND','tesseract'))
                    if command:
                        import pypdfium2
                        from agent_execution.process import run_process
                        with pypdfium2.PdfDocument(str(path)) as pdf,tempfile.TemporaryDirectory() as directory:
                            pdf_page=pdf[number-1]
                            width,height=pdf_page.get_size();scale=min(1.5,(8000000/max(1,width*height))**.5)
                            bitmap=pdf_page.render(scale=scale);image=bitmap.to_pil();target=Path(directory)/'page.png';image.save(target)
                            bitmap.close();pdf_page.close()
                            result=run_process([command,str(target),'stdout','-l',os.environ.get('AGENT_TESSERACT_LANGUAGE','eng')],timeout=10,max_output_bytes=131072)
                            if result.returncode==0:text=result.stdout.strip();ocr=True
                            else:needs_ocr=True
                    else:needs_ocr=True
                add('page',number,text,f'page:{number}',ocr=ocr,needs_ocr=needs_ocr)
            if len(document.pages)>200:truncated=True
    else:
        from defusedxml.ElementTree import fromstring
        with zipfile.ZipFile(path) as archive:
            entries=archive.infolist()
            if len(entries)>2048 or sum(v.file_size for v in entries)>32*1024*1024 or any(v.file_size>8*1024*1024 or v.flag_bits&1 for v in entries):raise ValueError('Office archive capacity/encryption limit')
            def xml(name):return fromstring(archive.read(name))
            if suffix=='.docx':
                root=xml('word/document.xml');namespace={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                for number,paragraph in enumerate(root.findall('.//w:p',namespace),1):
                    text=''.join(node.text or '' for node in paragraph.findall('.//w:t',namespace))
                    if text:add('paragraph',number,text,f'paragraph:{number}')
            elif suffix=='.pptx':
                import re
                names=sorted((v.filename for v in entries if re.fullmatch(r'ppt/slides/slide\d+\.xml',v.filename)),key=lambda name:int(re.search(r'(\d+)\.xml',name)[1]))
                for name in names:
                    number=int(re.search(r'(\d+)\.xml',name)[1]);root=xml(name)
                    text='\n'.join(node.text or '' for node in root.iter() if node.tag.endswith('}t'))
                    add('slide',number,text,f'slide:{number}')
            elif suffix=='.xlsx':
                from openpyxl import load_workbook
                workbook=load_workbook(path,read_only=True,data_only=True,keep_links=False)
                try:
                    for sheet in workbook:
                        for number,row in enumerate(sheet.iter_rows(values_only=True),1):
                            if len(units)>=200 or remaining<=0:truncated=True;break
                            if any(value is not None for value in row):add('row',number,' | '.join(str(value) if value is not None else '' for value in row),f'sheet:{sheet.title}/row:{number}')
                finally:workbook.close()
            else:raise ValueError('Unsupported document suffix')
    return {'success':True,'units':units,'truncated':truncated,'requires_ocr':any(unit.get('needs_ocr') for unit in units)}


if __name__=='__main__':
    try:result=parse(Path(sys.argv[1]))
    except Exception as error:result={'success':False,'error':type(error).__name__,'units':[]}
    Path(sys.argv[2]).write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
